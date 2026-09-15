from flask import Blueprint, request, jsonify
from db import execute_query
import csv
import io

employees_bp = Blueprint('employees', __name__)

@employees_bp.route('/api/employees', methods=['GET'])
def get_employees():
    department = request.args.get('department', '')
    search = request.args.get('search', '')
    
    query = "SELECT * FROM employees WHERE 1=1"
    params = []
    
    if department:
        query += " AND department = %s"
        params.append(department)
    if search:
        query += " AND (full_name LIKE %s OR emp_code LIKE %s OR email LIKE %s)"
        params.extend([f'%{search}%'] * 3)
    
    query += " ORDER BY employee_id"
    employees = execute_query(query, params)
    
    # Convert dates to strings
    for emp in employees:
        if emp.get('date_of_joining'):
            emp['date_of_joining'] = str(emp['date_of_joining'])
        if emp.get('created_at'):
            emp['created_at'] = str(emp['created_at'])
        if emp.get('updated_at'):
            emp['updated_at'] = str(emp['updated_at'])
    
    return jsonify(employees)

@employees_bp.route('/api/employees/<int:employee_id>', methods=['GET'])
def get_employee(employee_id):
    employees = execute_query("SELECT * FROM employees WHERE employee_id = %s", (employee_id,))
    if not employees:
        return jsonify({'error': 'Employee not found'}), 404
    emp = employees[0]
    if emp.get('date_of_joining'):
        emp['date_of_joining'] = str(emp['date_of_joining'])
    if emp.get('created_at'):
        emp['created_at'] = str(emp['created_at'])
    if emp.get('updated_at'):
        emp['updated_at'] = str(emp['updated_at'])
    return jsonify(emp)

@employees_bp.route('/api/employees', methods=['POST'])
def add_employee():
    data = request.json
    query = """INSERT INTO employees (emp_code, full_name, email, phone, department, designation, date_of_joining, employment_type, status)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)"""
    params = (
        data['emp_code'], data['full_name'], data.get('email', ''),
        data.get('phone', ''), data.get('department', ''), data.get('designation', ''),
        data.get('date_of_joining'), data.get('employment_type', 'full-time'),
        data.get('status', 'active')
    )
    employee_id = execute_query(query, params, fetch=False)
    
    # Also create salary structure if provided
    if data.get('basic_salary'):
        sal_query = """INSERT INTO salary_structure (employee_id, basic_salary, hra, da, travel_allowance, medical_allowance, pf_percent, tax_percent)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s)"""
        sal_params = (
            employee_id, data.get('basic_salary', 0), data.get('hra', 0),
            data.get('da', 0), data.get('travel_allowance', 0), data.get('medical_allowance', 0),
            data.get('pf_percent', 12), data.get('tax_percent', 5)
        )
        execute_query(sal_query, sal_params, fetch=False)
    
    return jsonify({'success': True, 'employee_id': employee_id}), 201

@employees_bp.route('/api/employees/<int:employee_id>', methods=['PUT'])
def update_employee(employee_id):
    data = request.json
    query = """UPDATE employees SET full_name=%s, email=%s, phone=%s, department=%s, 
               designation=%s, date_of_joining=%s, employment_type=%s, status=%s WHERE employee_id=%s"""
    params = (
        data['full_name'], data.get('email', ''), data.get('phone', ''),
        data.get('department', ''), data.get('designation', ''),
        data.get('date_of_joining'), data.get('employment_type', 'full-time'),
        data.get('status', 'active'), employee_id
    )
    execute_query(query, params, fetch=False)
    return jsonify({'success': True})

@employees_bp.route('/api/employees/<int:employee_id>', methods=['DELETE'])
def delete_employee(employee_id):
    execute_query("DELETE FROM employees WHERE employee_id = %s", (employee_id,), fetch=False)
    return jsonify({'success': True})

@employees_bp.route('/api/employees/upload', methods=['POST'])
def upload_employees():
    if 'file' not in request.files:
        return jsonify({'error': 'No file uploaded'}), 400
    
    file = request.files['file']
    filename = file.filename.lower()
    rows = []

    if filename.endswith('.csv'):
        stream = io.StringIO(file.stream.read().decode('utf-8', errors='replace'))
        reader = csv.DictReader(stream)
        rows = list(reader)
    elif filename.endswith(('.xlsx', '.xls')):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            headers = [str(c.value).strip() if c.value else '' for c in next(ws.iter_rows(min_row=1, max_row=1))]
            for row in ws.iter_rows(min_row=2, values_only=True):
                rows.append(dict(zip(headers, [str(v) if v is not None else '' for v in row])))
        except Exception as e:
            return jsonify({'error': f'Failed to read Excel file: {str(e)}'}), 400
    else:
        return jsonify({'error': 'Only CSV and Excel files supported'}), 400

    count = 0
    for row in rows:
        try:
            query = """INSERT INTO employees (emp_code, full_name, email, phone, department, designation, date_of_joining, employment_type)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s)"""
            params = (
                row.get('emp_code', ''), row.get('full_name', ''),
                row.get('email', ''), row.get('phone', ''),
                row.get('department', ''), row.get('designation', ''),
                row.get('date_of_joining') or None, row.get('employment_type', 'full-time')
            )
            execute_query(query, params, fetch=False)
            count += 1
        except Exception:
            continue
    
    return jsonify({'success': True, 'imported': count})

@employees_bp.route('/api/departments', methods=['GET'])
def get_departments():
    depts = execute_query("SELECT DISTINCT department FROM employees WHERE department IS NOT NULL AND department != '' ORDER BY department")
    return jsonify([d['department'] for d in depts])
